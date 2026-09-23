import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

from config import BASE_DIR, ENUMS_PATH, NATIVES_DIR
from src.database.missions.build import (
    MISSION_COLUMNS,
    MissionCatalog,
    QuestRecord,
    STREAM_TEXT_FIELDS,
    build_mission_workbook_data,
    load_mission_catalog,
)
from src.shared.text.catalog import TextSource, discover_language_ids
from src.database.missions.excel import style_mission_workbook, write_mission_workbook
from src.database.missions.health import calculate_solo_health
from src.shared.excel.cells import safe_cell


def _entry(guid, name, english, chinese=""):
    contents = [""] * 33
    contents[1] = english
    contents[13] = chinese
    return guid, name, contents


def _quest(target_type="EM_BOSS_HUNTING"):
    return {
        "_MissionId": "[123] MISSION_000123",
        "_Version": "[0] VER_1000_00_00",
        "_QuestType": "[0] HUNTING",
        "_QuestAttribute": "[1] SIDE_MISSION",
        "_QuestLv": 3,
        "_ClearCondition": {
            "app.user_data.QuestData.cClearCondition": {
                "_TargetType": f"[1] {target_type}",
                "_TargetInfoArray": [
                    {
                        "app.user_data.QuestData.cClearCondition.cTargetInfo": {
                            "_LegendaryID": "[0] NONE",
                            "_TargetIDValue": fixed_id,
                        }
                    }
                    for fixed_id in (7, 9)
                ],
            }
        },
        "_OrderCondition": {
            "app.user_data.QuestData.cOrderCondition": {
                "_MaxPlayerNum": 4,
                "_OrderHR": 3,
                "_OrderMR": 0,
            }
        },
        "_QuestMsg": {
            "app.user_data.QuestData.cQuestMsg": {
                "_TitleMsg": "title-guid",
                "_DetailMsg": "detail-guid",
            }
        },
        "_SubBossInfoArray": [
            {"app.user_data.QuestData.cSubBossInfo": {"_EmID": f"[0] {enemy_id}"}}
            for enemy_id in ("EM_A", "EM_C", "EM_A", "EM_B")
        ],
        "_TimeLimit": 50,
        "_RemMoney": 100,
        "_HRPoint": 10,
        "_AddPoint": 5,
        "_EnableGuestNpc": False,
        "_Stage": "[0] ST101",
        "_BattleBGM": "[0] DEFAULT",
        "_ClearBGM": "[0] DEFAULT",
    }


def _catalog(quest, stream_text=None, missions_without_quest_data=()):
    return MissionCatalog(
        (QuestRecord(Path("fixture_QuestData.user.3.json"), quest, stream_text),),
        {7: "EM_A", 9: "EM_B"},
        {"EM_A": "monster-guid", "EM_B": "missing-guid", "EM_C": "monster-c-guid"},
        missions_without_quest_data,
    )


def _text_source():
    return TextSource(
        [
            _entry("title-guid", None, "Hunt <EMID EM_A> and <EMID EM_MISSING>"),
            _entry("detail-guid", None, "English detail\r\nSecond line"),
            _entry("monster-guid", "EnemyText_NAME_EM_A", "Monster", "怪物"),
            _entry("monster-c-guid", "EnemyText_NAME_EM_C", "Monster", "怪物"),
            _entry(None, "MsgGUI020203_0001", "Side Mission"),
        ]
    )


class MissionFixtureTests(unittest.TestCase):
    def test_numeric_cells_are_centered_without_centering_numeric_text_or_lists(self):
        data = build_mission_workbook_data(_catalog(_quest()), _text_source(), 13)
        data.rows[0]["_RemMoney"] = 1.25
        data.rows[0]["_Stage"] = "123"
        data.rows[0]["_SubBossInfoArray"] = "100、200"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "MissionData"
        sheet.append([None] * len(data.headers))
        sheet.append([None] * len(data.headers))
        for row in data.rows:
            sheet.append([safe_cell(row[header.key]) for header in data.headers])
        style_mission_workbook(workbook, data)
        try:
            for column, expected in (
                ("_QuestLv", "center"),
                ("_RemMoney", "center"),
                ("_EnableGuestNpc", "left"),
                ("_Stage", "left"),
                ("_SubBossInfoArray", "left"),
            ):
                cell = sheet.cell(3, data.columns.index(column) + 1)
                self.assertEqual((cell.alignment.horizontal, cell.alignment.vertical),
                                 (expected, "center"), column)
        finally:
            workbook.close()

    def test_solo_health_uses_new_random_grades_from_user3(self):
        self.assertEqual(calculate_solo_health(
            enemy_id="EM9999_00_0", base_health=10000, quest_health_rate=1.0,
            legendary_id="NORMAL", reward_rank=0, legendary_rates={"HealthRate": 1.0},
            random_rate_table={"_Value0": 1.0, "_ValueP4": 1.125},
            random_probability_tables=[{"_Prob0": 50, "_ProbP4": 50}],
            difficulty_adjust_range=0, king_when_none_ids=frozenset(),
            no_auto_hard_ids=frozenset(),
        ), "10000、11250")

    def test_solo_health_compiled_exceptions_are_explicit_inputs(self):
        inputs = {
            "enemy_id": "EM9999_00_0", "base_health": 10000,
            "quest_health_rate": 1.0, "reward_rank": 10,
            "legendary_rates": {
                "HealthRate": 1.0, "HealthRate_Hard": 1.5,
                "HealthRate_King": 1.0, "HealthRate_King_Hard": 1.25,
            },
            "random_rate_table": {"_Value0": 1.0},
            "random_probability_tables": [{"_Prob0": 100}],
            "difficulty_adjust_range": 0,
        }
        self.assertEqual(calculate_solo_health(
            **inputs, legendary_id="NORMAL", king_when_none_ids=frozenset(),
            no_auto_hard_ids=frozenset({"EM9999_00_0"}),
        ), 10000)
        self.assertEqual(calculate_solo_health(
            **inputs, legendary_id="NORMAL", king_when_none_ids=frozenset(),
            no_auto_hard_ids=frozenset({"EM9999_00_0"}), preset_hard=True,
        ), 15000)
        self.assertEqual(calculate_solo_health(
            **inputs, legendary_id="NONE",
            king_when_none_ids=frozenset({"EM9999_00_0"}),
            no_auto_hard_ids=frozenset(),
        ), 12500)
        with self.assertRaisesRegex(ValueError, "Ambiguous rank-10 NONE"):
            calculate_solo_health(
                **inputs, legendary_id="NONE", king_when_none_ids=frozenset(),
                no_auto_hard_ids=frozenset(),
            )

    def test_msdata_id_without_quest_data_is_appended_with_available_text(self):
        source = TextSource(_text_source().entries + [
            _entry(None, "Mission000124_100", "Quest title"),
            _entry(None, "Mission000124_102", "Quest detail"),
            _entry(None, "Mission000124_000", "Story title", "故事标题"),
            _entry(None, "Mission000125_000", "Story title", "故事标题"),
        ])
        missing = ("MISSION_000124", "MISSION_000125", "MISSION_000126")
        data = build_mission_workbook_data(
            _catalog(_quest(), missions_without_quest_data=missing), source, 13
        )
        self.assertEqual(data.groups, ((3, 4), (5, 5), (6, 6), (7, 7)))
        self.assertEqual([row["_MissionId"] for row in data.rows[-3:]], list(missing))
        self.assertEqual(data.rows[-3]["_TitleMsg"], "Quest title")
        self.assertEqual(data.rows[-3]["_DetailMsg"], "Quest detail")
        self.assertEqual(data.rows[-2]["_TitleMsg"], "故事标题")
        self.assertTrue(all(
            row[column] == ""
            for row in data.rows[-3:]
            for column in MISSION_COLUMNS
            if column not in ("_MissionId", "_TitleMsg", "_DetailMsg")
        ))
        self.assertEqual(data.rows[-1]["_TitleMsg"], "")

    def test_language_fallback_monster_id_and_deduplicated_subbosses(self):
        data = build_mission_workbook_data(_catalog(_quest()), _text_source(), 13)
        self.assertEqual(len(data.rows), 2)
        self.assertEqual(data.rows[0]["_QuestType"], "HUNTING")
        self.assertEqual(data.rows[0]["_QuestAttribute"], "Side Mission")
        self.assertEqual(data.rows[0]["_TitleMsg"], "Hunt 怪物 and EM_MISSING")
        self.assertEqual(data.rows[0]["_DetailMsg"], "English detail\nSecond line")
        self.assertEqual([row["MonsterName"] for row in data.rows], ["怪物", "EM_B"])
        self.assertEqual(data.rows[0]["_SubBossInfoArray"], "怪物\nEM_B")

    def test_monster_name_uses_another_language_before_em_id(self):
        entries = _text_source().entries.copy()
        guid, name, contents = entries[2]
        contents = contents.copy()
        contents[1] = ""
        entries[2] = (guid, name, contents)
        data = build_mission_workbook_data(_catalog(_quest()), TextSource(entries), 0)
        self.assertEqual(data.rows[0]["MonsterName"], "怪物")
        self.assertEqual(data.rows[0]["_TitleMsg"], "Hunt 怪物 and EM_MISSING")

    def test_enemy_ref_uses_em_id_when_all_names_are_missing(self):
        entries = _text_source().entries.copy()
        guid, name, contents = entries[0]
        contents = contents.copy()
        contents[1] = "Find <REF EnemyText_NAME_EM_A> and <REF EnemyText_NAME_EM_MISSING>"
        entries[0] = (guid, name, contents)
        data = build_mission_workbook_data(_catalog(_quest()), TextSource(entries), 13)
        self.assertEqual(data.rows[0]["_TitleMsg"], "Find 怪物 and EM_MISSING")

    def test_named_monster_text_is_used_when_enemy_guid_is_absent(self):
        catalog = _catalog(_quest())
        catalog.enemy_name_guids.clear()
        data = build_mission_workbook_data(catalog, _text_source(), 13)
        self.assertEqual([row["MonsterName"] for row in data.rows], ["怪物", "EM_B"])

    def test_item_condition_leaves_target_monster_fields_empty(self):
        data = build_mission_workbook_data(_catalog(_quest("ITEM")), _text_source(), 13)
        self.assertEqual(len(data.rows), 1)
        self.assertEqual(data.rows[0]["_TargetType"], "ITEM")
        self.assertTrue(all(not row["_LegendaryID"] and not row["MonsterName"] for row in data.rows))

    def test_multiple_clear_conditions_use_consecutive_rows(self):
        quest = _quest()
        original = quest["_ClearCondition"]
        quest["_ClearCondition"] = [
            original,
            {
                "app.user_data.QuestData.cClearCondition": {
                    "_TargetType": "[6] ITEM",
                    "_TargetInfoArray": [],
                }
            },
        ]
        data = build_mission_workbook_data(_catalog(quest), _text_source(), 13)
        self.assertEqual(data.groups, ((3, 5),))
        self.assertEqual([row["_TargetType"] for row in data.rows], [
            "EM_BOSS_HUNTING", "EM_BOSS_HUNTING", "ITEM"
        ])
        self.assertEqual(data.rows[-1]["MonsterName"], "")

    def test_stream_text_uses_sibling_data_with_english_fallback(self):
        stream = {
            "_IDList": ["Mission000123_100", "Mission000123_102"],
            "_EnglishList": ["Stream title", "English stream detail"],
            "_SimplelifiedChineseList": ["", "中文说明"],
        }
        data = build_mission_workbook_data(_catalog(_quest(), stream), _text_source(), 13)
        self.assertEqual(data.rows[0]["_TitleMsg"], "Stream title")
        self.assertEqual(data.rows[0]["_DetailMsg"], "中文说明")


class MissionCorpusTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.catalog = load_mission_catalog(NATIVES_DIR, ENUMS_PATH)
        cls.text_source = TextSource.from_natives(NATIVES_DIR)

    def test_every_export_language_has_complete_quest_rows(self):
        languages = discover_language_ids(NATIVES_DIR)
        self.assertEqual(set(languages), set(STREAM_TEXT_FIELDS))
        expected_rows = sum(
            max(1, len(record.data["_ClearCondition"]["app.user_data.QuestData.cClearCondition"]["_TargetInfoArray"]))
            for record in self.catalog.quests
        )
        self.assertTrue(any(record.stream_text for record in self.catalog.quests))
        self.assertGreater(expected_rows, len(self.catalog.quests))
        self.assertIn("MISSION_400063", self.catalog.missions_without_quest_data)
        self.assertNotIn("MISSION_001020", self.catalog.missions_without_quest_data)
        for language_id in languages:
            with self.subTest(language_id=language_id):
                data = build_mission_workbook_data(self.catalog, self.text_source, language_id)
                self.assertEqual(
                    len(data.groups),
                    len(self.catalog.quests) + len(self.catalog.missions_without_quest_data),
                )
                self.assertEqual(len(data.rows), expected_rows + len(self.catalog.missions_without_quest_data))
                self.assertTrue(all(row["_TitleMsg"] and row["_DetailMsg"] for row in data.rows[:expected_rows]))
                self.assertEqual(
                    [row["_MissionId"] for row in data.rows[expected_rows:]],
                    list(self.catalog.missions_without_quest_data),
                )
                self.assertTrue(all(
                    row["MonsterName"]
                    for row in data.rows[:expected_rows] if row["_TargetType"].startswith("EM_")
                ))
                self.assertTrue(all(
                    len(names := row["_SubBossInfoArray"].split("\n")) == len(set(names))
                    for row in data.rows if row["_SubBossInfoArray"]
                ))

    def test_msdata_only_examples_fill_only_confirmed_columns(self):
        data = build_mission_workbook_data(self.catalog, self.text_source, 1)
        rows = {row["_MissionId"]: row for row in data.rows}
        self.assertEqual(rows["MISSION_001040"]["_TitleMsg"], "Back to Camp")
        self.assertTrue(rows["MISSION_001040"]["_DetailMsg"])
        self.assertTrue(all(
            value == "" for column, value in rows["MISSION_400063"].items()
            if column != "_MissionId"
        ))

    def test_target_rates_resolve_selected_layouts_and_stream_quests(self):
        difficulty = self.catalog.difficulty
        self.assertIsNotNone(difficulty)
        self.assertEqual(len(difficulty.targets), 241)
        self.assertEqual(difficulty.schema.counts, (2, 3, 4, 5))
        self.assertEqual(len(difficulty.schema.statuses), 9)

        data = build_mission_workbook_data(self.catalog, self.text_source, 1)
        rows = {row["_MissionId"]: row for row in data.rows}
        self.assertEqual(data.columns[:9], (
            "_MissionId", "_QuestLv", "_OrderHR", "_OrderMR", "_TitleMsg",
            "MonsterName", "_TargetType", "_LegendaryID", "_RewardRank",
        ))
        self.assertEqual(len(data.columns), 108)
        self.assertEqual(data.columns.index("SoloHealth") + 1, data.columns.index("_Health"))
        self.assertEqual(rows["MISSION_001130"]["_Health"], 1.17)
        self.assertEqual(rows["MISSION_001130"]["_Attack"], 1.0)
        self.assertEqual(rows["MISSION_001130"]["_Count=2._Health"], 1.6)
        self.assertEqual(
            rows["MISSION_004320"]["_MultiTableId"],
            "15789b29-d7df-4354-8e1e-64999b31ab49",
        )
        self.assertEqual(
            rows["MISSION_600000"]["_DifficultyRankId"],
            rows["MISSION_001130"]["_DifficultyRankId"],
        )
        for mission_id in ("MISSION_101001", "MISSION_640001", "MISSION_400063"):
            self.assertEqual(rows[mission_id]["SoloHealth"], "")
            self.assertEqual(rows[mission_id]["_Health"], "")
            self.assertEqual(rows[mission_id]["_Count=2._Health"], "")

    def test_solo_health_matches_observed_mission_values(self):
        data = build_mission_workbook_data(self.catalog, self.text_source, 1)
        rows = {row["_MissionId"]: row for row in data.rows}
        for mission_id, expected in (
            ("MISSION_109011", 47840),
            ("MISSION_661002", 28944),
            ("MISSION_650006", 78566),
            ("MISSION_204000", 6000),
            ("MISSION_650000", "76726、78325、79923、81522、83120"),
        ):
            with self.subTest(mission_id=mission_id):
                self.assertEqual(rows[mission_id]["SoloHealth"], expected)

    def test_workbook_merges_mission_fields_and_centers_numeric_content(self):
        data = build_mission_workbook_data(self.catalog, self.text_source, 1)
        BASE_DIR.joinpath(".agents").mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=BASE_DIR / ".agents") as temp_dir:
            path = write_mission_workbook(Path(temp_dir) / "MissionData.xlsx", data)
            workbook = load_workbook(path)
            try:
                sheet = workbook["MissionData"]
                self.assertEqual(sheet.max_column, len(data.columns))
                self.assertEqual(sheet.max_row, len(data.rows) + 2)
                self.assertEqual(sheet.freeze_panes, "I3")
                solo = data.columns.index("SoloHealth") + 1
                self.assertEqual(sheet.cell(1, solo).value, "SoloHealth")
                self.assertEqual(sheet.cell(1, solo + 1).value, "_Health")
                self.assertEqual([sheet.cell(1, index).value for index in range(1, 10)], [
                    "_MissionId", "_QuestLv", "_OrderHR", "_OrderMR", "_TitleMsg",
                    "MonsterName", "_TargetType", "_LegendaryID", "_RewardRank",
                ])
                merged = {str(region) for region in sheet.merged_cells.ranges}
                self.assertIn("A1:A2", merged)
                poison = data.columns.index("_Poison._DefaultLimit") + 1
                poison_next = get_column_letter(poison + 1)
                poison_letter = get_column_letter(poison)
                self.assertIn(f"{poison_letter}1:{poison_next}1", merged)
                self.assertEqual(sheet.cell(1, poison).value, "_Poison")
                self.assertEqual(sheet.cell(2, poison).value, "_DefaultLimit")
                self.assertEqual(sheet.cell(2, poison + 1).value, "_AddAndMaxLimit")
                count = data.columns.index("_Count=2._Health") + 1
                self.assertIn(
                    f"{get_column_letter(count)}1:{get_column_letter(count + 12)}1", merged
                )
                self.assertEqual(sheet.cell(1, count).value, "_Count = 2")
                self.assertEqual(sheet.cell(2, count).value, "_Health")
                self.assertNotEqual(
                    sheet.cell(1, 1).fill.fgColor.rgb,
                    sheet.cell(1, 9).fill.fgColor.rgb,
                )
                self.assertNotEqual(
                    sheet.cell(3, 1).fill.fgColor.rgb,
                    sheet.cell(3, 9).fill.fgColor.rgb,
                )
                first, last = next(group for group in data.groups if group[1] > group[0])
                for column in (
                    "_MissionId", "_QuestLv", "_OrderHR", "_OrderMR", "_TitleMsg",
                    "_DetailMsg", "_SubBossInfoArray",
                ):
                    letter = get_column_letter(data.columns.index(column) + 1)
                    self.assertIn(f"{letter}{first}:{letter}{last}", merged)
                for column in ("MonsterName", "_TargetType", "_Health", "_Count=2._Health"):
                    letter = get_column_letter(data.columns.index(column) + 1)
                    self.assertNotIn(f"{letter}{first}:{letter}{last}", merged)
                for row in sheet.iter_rows(min_row=3):
                    for cell in row:
                        if isinstance(cell, MergedCell):
                            continue
                        numeric = isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
                        self.assertEqual(
                            (cell.alignment.horizontal, cell.alignment.vertical),
                            ("center" if numeric else "left", "center"),
                            cell.coordinate,
                        )
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
