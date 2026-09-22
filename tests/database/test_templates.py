from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from src.database.table import TableWorkbook
from src.shared.equipment import prepare_equipment
from src.shared.text.catalog import TextDB
from src.shared.text.values import TextRef


class TemplateTests(unittest.TestCase):
    def test_compound_skill_names_are_independent_of_language_order(self):
        skill_guid = "11111111-1111-1111-1111-111111111111"
        name_guid = "22222222-2222-2222-2222-222222222222"
        tables = prepare_equipment({"Armor": [{
            "Name": TextRef(name_guid), "SlotLevel": [1], "Skill": ["S001"], "SkillLevel": [2],
        }]}, lambda path: [{"skillId": "S001", "skillName": TextRef(skill_guid)}]
            if path.endswith("SkillCommonData.user.3.json") else [])
        template = TableWorkbook("Equipment.xlsx", tables)
        languages = {
            "en": TextDB({name_guid: "Armor", skill_guid: "Attack"}, {}),
            "zh": TextDB({name_guid: "防具", skill_guid: "攻击"}, {}),
        }
        snapshots = {}
        with tempfile.TemporaryDirectory() as directory:
            for order in (tuple(languages), tuple(reversed(languages))):
                for language in order:
                    path = template.write(Path(directory) / language, languages[language])
                    workbook = load_workbook(path)
                    values = list(workbook.active.values)
                    workbook.close()
                    if language in snapshots:
                        self.assertEqual(snapshots[language], values)
                    snapshots[language] = values
        self.assertIn("['Attack: 2']", snapshots["en"][1])
        self.assertIn("['攻击: 2']", snapshots["zh"][1])
        self.assertIsInstance(template.sheets["Armor"][0]["Name"], TextRef)
