from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from src.database.full_text.exporter import export_full_text
from src.shared.excel.writer import write_workbook
from src.shared.rarity import Rarity, prepare_rarities
from src.shared.text.catalog import TextDB


class ExcelTests(unittest.TestCase):
    def test_rarity_is_explicit_and_plain_text_is_not_reinterpreted(self):
        prepared = prepare_rarities({"Data": [{"Rare": "RARE7", "Explain": "Text"}]})
        self.assertEqual(prepared["Data"][0]["Rare"], Rarity(7))
        with tempfile.TemporaryDirectory() as directory:
            path = write_workbook(Path(directory) / "data.xlsx", prepared, 80)
            workbook = load_workbook(path)
            self.assertEqual(workbook.active["A2"].value, 8)
            self.assertEqual(workbook.active["A2"].fill.fill_type, "solid")
            workbook.close()
            full_text = export_full_text(Path(directory), TextDB({"guid": "RARE7"}, {}))
            workbook = load_workbook(full_text)
            self.assertEqual(workbook.active["B2"].value, "RARE7")
            workbook.close()
        self.assertEqual(prepared["Data"][0]["Rare"], Rarity(7))

    def test_full_text_escapes_formulas_and_handles_an_empty_catalog(self):
        with tempfile.TemporaryDirectory() as directory:
            path = export_full_text(Path(directory), TextDB({"guid": "=1+1"}, {}))
            workbook = load_workbook(path)
            self.assertEqual(workbook.active["B2"].value, "'=1+1")
            workbook.close()
            path = export_full_text(Path(directory), TextDB({}, {}))
            workbook = load_workbook(path)
            self.assertEqual(list(workbook.active.values), [("guid", "text")])
            workbook.close()
