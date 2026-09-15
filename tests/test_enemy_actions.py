import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from config import SUPPORT_FILES
from src.converters.enemy_actions import (
    INDEX_SHEET_NAME,
    load_enemy_action_catalog,
    write_enemy_action_workbook,
)
from src.data.text_db import TextDB


ENEMY_NAME_GUID_1 = "00000000-0000-0000-0000-000000000001"
ENEMY_NAME_GUID_2 = "00000000-0000-0000-0000-000000000002"


class EnemyActionTests(unittest.TestCase):
    def test_catalog_discovers_filename_variant_and_stably_deduplicates_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            natives_dir = Path(temp_dir)
            _write_enemy_data(
                natives_dir,
                [
                    ("EM0001_00_0", ENEMY_NAME_GUID_1),
                    ("EM5001_10_0", ENEMY_NAME_GUID_2),
                ],
            )
            _write_shell_creator(
                natives_dir,
                "Em0001",
                "00",
                "Em0001_00_ShellCreatorInfo.user.3.json",
                [
                    ("Attack", "Comment"),
                    ("Attack", "Comment"),
                    (" Name only ", ""),
                    ("", " Comment only "),
                    ("", ""),
                ],
            )
            _write_shell_creator(
                natives_dir,
                "Em5001",
                "10",
                "Em5001_10_ShellCreatorInfoData.user.3.json",
                [("ExplosionLiquid", ""), ("ExplosionSpark", "")],
            )
            text_db = TextDB(
                {
                    ENEMY_NAME_GUID_1: "怪物一",
                    ENEMY_NAME_GUID_2: "爆炸瓦斯蛙",
                },
                {},
            )

            catalog = load_enemy_action_catalog(natives_dir, text_db)

        self.assertEqual([sheet.enemy_id for sheet in catalog.sheets], ["EM0001_00_0", "EM5001_10_0"])
        self.assertEqual(
            catalog.sheets[0].rows,
            (
                ("Attack", "Comment"),
                ("Name only", ""),
                ("", "Comment only"),
            ),
        )
        self.assertEqual(catalog.raw_rows, 7)
        self.assertEqual(catalog.displayed_rows, 5)
        self.assertEqual(catalog.duplicate_rows, 1)
        self.assertEqual(catalog.blank_rows, 1)

    def test_full_enemy_id_mapping_must_be_unique(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            natives_dir = Path(temp_dir)
            _write_enemy_data(
                natives_dir,
                [
                    ("EM1150_00_0", ENEMY_NAME_GUID_1),
                    ("EM1150_00_1", ENEMY_NAME_GUID_2),
                ],
            )
            _write_shell_creator(
                natives_dir,
                "Em1150",
                "00",
                "Em1150_00_ShellCreatorInfo.user.3.json",
                [("Attack", "Comment")],
            )
            text_db = TextDB(
                {
                    ENEMY_NAME_GUID_1: "雷甲龙♂",
                    ENEMY_NAME_GUID_2: "雷甲龙♀",
                },
                {},
            )

            with self.assertRaisesRegex(ValueError, "exactly one full enemy ID"):
                load_enemy_action_catalog(natives_dir, text_db)

    def test_workbook_has_index_links_and_one_two_column_sheet_per_enemy(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            natives_dir = root / "natives"
            _write_enemy_data(
                natives_dir,
                [
                    ("EM0001_00_0", ENEMY_NAME_GUID_1),
                    ("EM5001_10_0", ENEMY_NAME_GUID_2),
                ],
            )
            _write_shell_creator(
                natives_dir,
                "Em0001",
                "00",
                "Em0001_00_ShellCreatorInfo.user.3.json",
                [("Attack", "Comment"), ("Attack", "Comment")],
            )
            _write_shell_creator(
                natives_dir,
                "Em5001",
                "10",
                "Em5001_10_ShellCreatorInfoData.user.3.json",
                [("ExplosionLiquid", "")],
            )
            text_db = TextDB(
                {
                    ENEMY_NAME_GUID_1: "怪物一",
                    ENEMY_NAME_GUID_2: "爆炸瓦斯蛙",
                },
                {},
            )
            catalog = load_enemy_action_catalog(natives_dir, text_db)
            path = write_enemy_action_workbook(root / "EnemyActionNames.xlsx", catalog)
            workbook = load_workbook(path, read_only=False, data_only=False)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    [
                        INDEX_SHEET_NAME,
                        "怪物一（EM0001_00_0）",
                        "爆炸瓦斯蛙（EM5001_10_0）",
                    ],
                )
                index = workbook[INDEX_SHEET_NAME]
                self.assertEqual(index["A1"].value, "怪物EM编号")
                self.assertEqual(index["B1"].value, "怪物名称")
                self.assertEqual(index["A2"].value, "EM0001_00_0")
                self.assertEqual(index["B2"].value, "怪物一")
                self.assertEqual(
                    index["B2"].hyperlink.location,
                    "'怪物一（EM0001_00_0）'!A1",
                )
                self.assertEqual(index.freeze_panes, "A2")
                self.assertEqual(index.auto_filter.ref, "A1:B3")

                action_sheet = workbook["怪物一（EM0001_00_0）"]
                self.assertEqual(action_sheet.max_row, 2)
                self.assertEqual(action_sheet.max_column, 2)
                self.assertEqual(action_sheet["A1"].value, "name")
                self.assertEqual(action_sheet["B1"].value, "comment")
                self.assertEqual(action_sheet["A2"].value, "Attack")
                self.assertEqual(action_sheet["B2"].value, "Comment")
                self.assertEqual(action_sheet.freeze_panes, "A2")
                self.assertEqual(action_sheet.auto_filter.ref, "A1:B2")
                self.assertFalse(action_sheet.sheet_view.showGridLines)
            finally:
                workbook.close()


def _write_enemy_data(
    natives_dir: Path,
    enemies: list[tuple[str, str]],
) -> None:
    rows = []
    for index, (enemy_id, name_guid) in enumerate(enemies, start=1):
        rows.append(
            {
                "app.user_data.EnemyData.cData": {
                    "_Index": index,
                    "_enemyId": f"[{index}] {enemy_id}",
                    "_EnemyName": name_guid,
                }
            }
        )
    _write_json(
        natives_dir / SUPPORT_FILES["enemy"],
        [{"app.user_data.EnemyData": {"_Values": rows}}],
    )


def _write_shell_creator(
    natives_dir: Path,
    em_directory: str,
    variant: str,
    filename: str,
    rows: list[tuple[str, str]],
) -> None:
    records = []
    for index, (name, comment) in enumerate(rows):
        records.append(
            {
                "ace.user_data.ShellCreatorInfoData.ShellCreatorInfo": {
                    "_Name": name,
                    "_Comment": comment,
                    "_UniqueID": index,
                    "_ShellListNo": index,
                }
            }
        )
    _write_json(
        natives_dir
        / "STM/GameDesign/Enemy"
        / em_directory
        / variant
        / "Shell"
        / filename,
        [
            {
                "ace.user_data.ShellCreatorInfoData": {
                    "_ShellCreatorInfos": records,
                }
            }
        ],
    )


def _write_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    unittest.main()
