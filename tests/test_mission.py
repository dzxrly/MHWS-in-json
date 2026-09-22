import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from config import BASE_DIR, ENUMS_PATH, NATIVES_DIR
from src.converters.mission import (
    MISSION_COLUMNS,
    MissionCatalog,
    QuestRecord,
    STREAM_TEXT_FIELDS,
    build_mission_workbook_data,
    load_mission_catalog,
)
from src.data.text_db import TextSource, discover_language_ids
from src.excel.mission import write_mission_workbook


def _entry(guid, name, english, chinese=""):
    contents = [""] * 33
    contents[1] = english
    contents[13] = chinese
    return guid, name, contents


def _quest(target_type="EM_BOSS_HUNTING"):
    return {
        "_MissionId": "[123] MISSION_000123",
        "_Version": "[0] VER_1000_00_00",
        "_QuestType": "[0] HUNTING",
        "_QuestAttribute": "[1] SIDE_MISSION",
        "_QuestLv": 3,
        "_ClearCondition": {
            "app.user_data.QuestData.cClearCondition": {
                "_TargetType": f"[1] {target_type}",
                "_TargetInfoArray": [
                    {
                        "app.user_data.QuestData.cClearCondition.cTargetInfo": {
                            "_LegendaryID": "[0] NONE",
                            "_TargetIDValue": fixed_id,
                        }
                    }
                    for fixed_id in (7, 9)
                ],
            }
        },
        "_OrderCondition": {
            "app.user_data.QuestData.cOrderCondition": {
                "_MaxPlayerNum": 4,
                "_OrderHR": 3,
                "_OrderMR": 0,
            }
        },
        "_QuestMsg": {
            "app.user_data.QuestData.cQuestMsg": {
                "_TitleMsg": "title-guid",
                "_DetailMsg": "detail-guid",
            }
        },
        "_SubBossInfoArray": [
            {"app.user_data.QuestData.cSubBossInfo": {"_EmID": f"[0] {enemy_id}"}}
            for enemy_id in ("EM_A", "EM_C", "EM_A", "EM_B")
        ],
        "_TimeLimit": 50,
        "_RemMoney": 100,
        "_HRPoint": 10,
        "_AddPoint": 5,
        "_EnableGuestNpc": False,
        "_Stage": "[0] ST101",
        "_BattleBGM": "[0] DEFAULT",
        "_ClearBGM": "[0] DEFAULT",
    }


def _catalog(quest, stream_text=None, missions_without_quest_data=()):
    return MissionCatalog(
        (QuestRecord(Path("fixture_QuestData.user.3.json"), quest, stream_text),),
        {7: "EM_A", 9: "EM_B"},
        {"EM_A": "monster-guid", "EM_B": "missing-guid", "EM_C": "monster-c-guid"},
        missions_without_quest_data,
    )


def _text_source():
    return TextSource(
        [
            _entry("title-guid", None, "Hunt <EMID EM_A> and <EMID EM_MISSING>"),
            _entry("detail-guid", None, "English detail\r\nSecond line"),
            _entry("monster-guid", "EnemyText_NAME_EM_A", "Monster", "怪物"),
            _entry("monster-c-guid", "EnemyText_NAME_EM_C", "Monster", "怪物"),
            _entry(None, "MsgGUI020203_0001", "Side Mission"),
        ]
    )


class MissionFixtureTests(unittest.TestCase):
    def test_msdata_id_without_quest_data_is_appended_with_available_text(self):
        source = TextSource(_text_source().entries + [
            _entry(None, "Mission000124_100", "Quest title"),
            _entry(None, "Mission000124_102", "Quest detail"),
            _entry(None, "Mission000124_000", "Story title", "故事标题"),
            _entry(None, "Mission000125_000", "Story title", "故事标题"),
        ])
        missing = ("MISSION_000124", "MISSION_000125", "MISSION_000126")
        data = build_mission_workbook_data(
            _catalog(_quest(), missions_without_quest_data=missing), source, 13
        )
        self.assertEqual(data.groups, ((2, 3), (4, 4), (5, 5), (6, 6)))
        self.assertEqual([row["_MissionId"] for row in data.rows[-3:]], list(missing))
        self.assertEqual(data.rows[-3]["_TitleMsg"], "Quest title")
        self.assertEqual(data.rows[-3]["_DetailMsg"], "Quest detail")
        self.assertEqual(data.rows[-2]["_TitleMsg"], "故事标题")
        self.assertTrue(all(
            row[column] == ""
            for row in data.rows[-3:]
            for column in MISSION_COLUMNS
            if column not in ("_MissionId", "_TitleMsg", "_DetailMsg")
        ))
        self.assertEqual(data.rows[-1]["_TitleMsg"], "")

    def test_language_fallback_monster_id_and_deduplicated_subbosses(self):
        data = build_mission_workbook_data(_catalog(_quest()), _text_source(), 13)
        self.assertEqual(len(data.rows), 2)
        self.assertEqual(data.rows[0]["_QuestType"], "HUNTING")
        self.assertEqual(data.rows[0]["_QuestAttribute"], "Side Mission")
        self.assertEqual(data.rows[0]["_TitleMsg"], "Hunt 怪物 and EM_MISSING")
        self.assertEqual(data.rows[0]["_DetailMsg"], "English detail\nSecond line")
        self.assertEqual([row["MonsterName"] for row in data.rows], ["怪物", "EM_B"])
        self.assertEqual(data.rows[0]["_SubBossInfoArray"], "怪物\nEM_B")

    def test_monster_name_uses_another_language_before_em_id(self):
        entries = _text_source().entries.copy()
        guid, name, contents = entries[2]
        contents = contents.copy()
        contents[1] = ""
        entries[2] = (guid, name, contents)
        data = build_mission_workbook_data(_catalog(_quest()), TextSource(entries), 0)
        self.assertEqual(data.rows[0]["MonsterName"], "怪物")
        self.assertEqual(data.rows[0]["_TitleMsg"], "Hunt 怪物 and EM_MISSING")

    def test_enemy_ref_uses_em_id_when_all_names_are_missing(self):
        entries = _text_source().entries.copy()
        guid, name, contents = entries[0]
        contents = contents.copy()
        contents[1] = "Find <REF EnemyText_NAME_EM_A> and <REF EnemyText_NAME_EM_MISSING>"
        entries[0] = (guid, name, contents)
        data = build_mission_workbook_data(_catalog(_quest()), TextSource(entries), 13)
        self.assertEqual(data.rows[0]["_TitleMsg"], "Find 怪物 and EM_MISSING")

    def test_named_monster_text_is_used_when_enemy_guid_is_absent(self):
        catalog = _catalog(_quest())
        catalog.enemy_name_guids.clear()
        data = build_mission_workbook_data(catalog, _text_source(), 13)
        self.assertEqual([row["MonsterName"] for row in data.rows], ["怪物", "EM_B"])

    def test_item_condition_leaves_target_monster_fields_empty(self):
        data = build_mission_workbook_data(_catalog(_quest("ITEM")), _text_source(), 13)
        self.assertEqual(len(data.rows), 1)
        self.assertEqual(data.rows[0]["_TargetType"], "ITEM")
        self.assertTrue(all(not row["_LegendaryID"] and not row["MonsterName"] for row in data.rows))

    def test_multiple_clear_conditions_use_consecutive_rows(self):
        quest = _quest()
        original = quest["_ClearCondition"]
        quest["_ClearCondition"] = [
            original,
            {
                "app.user_data.QuestData.cClearCondition": {
                    "_TargetType": "[6] ITEM",
                    "_TargetInfoArray": [],
                }
            },
        ]
        data = build_mission_workbook_data(_catalog(quest), _text_source(), 13)
        self.assertEqual(data.groups, ((2, 4),))
        self.assertEqual([row["_TargetType"] for row in data.rows], [
            "EM_BOSS_HUNTING", "EM_BOSS_HUNTING", "ITEM"
        ])
        self.assertEqual(data.rows[-1]["MonsterName"], "")

    def test_stream_text_uses_sibling_data_with_english_fallback(self):
        stream = {
            "_IDList": ["Mission000123_100", "Mission000123_102"],
            "_EnglishList": ["Stream title", "English stream detail"],
            "_SimplelifiedChineseList": ["", "中文说明"],
        }
        data = build_mission_workbook_data(_catalog(_quest(), stream), _text_source(), 13)
        self.assertEqual(data.rows[0]["_TitleMsg"], "Stream title")
        self.assertEqual(data.rows[0]["_DetailMsg"], "中文说明")


class MissionCorpusTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.catalog = load_mission_catalog(NATIVES_DIR, ENUMS_PATH)
        cls.text_source = TextSource.from_natives(NATIVES_DIR)

    def test_every_export_language_has_complete_quest_rows(self):
        languages = discover_language_ids(NATIVES_DIR)
        self.assertEqual(set(languages), set(STREAM_TEXT_FIELDS))
        expected_rows = sum(
            max(1, len(record.data["_ClearCondition"]["app.user_data.QuestData.cClearCondition"]["_TargetInfoArray"]))
            for record in self.catalog.quests
        )
        self.assertTrue(any(record.stream_text for record in self.catalog.quests))
        self.assertGreater(expected_rows, len(self.catalog.quests))
        self.assertIn("MISSION_400063", self.catalog.missions_without_quest_data)
        self.assertNotIn("MISSION_001020", self.catalog.missions_without_quest_data)
        for language_id in languages:
            with self.subTest(language_id=language_id):
                data = build_mission_workbook_data(self.catalog, self.text_source, language_id)
                self.assertEqual(
                    len(data.groups),
                    len(self.catalog.quests) + len(self.catalog.missions_without_quest_data),
                )
                self.assertEqual(len(data.rows), expected_rows + len(self.catalog.missions_without_quest_data))
                self.assertTrue(all(row["_TitleMsg"] and row["_DetailMsg"] for row in data.rows[:expected_rows]))
                self.assertEqual(
                    [row["_MissionId"] for row in data.rows[expected_rows:]],
                    list(self.catalog.missions_without_quest_data),
                )
                self.assertTrue(all(
                    row["MonsterName"]
                    for row in data.rows[:expected_rows] if row["_TargetType"].startswith("EM_")
                ))
                self.assertTrue(all(
                    len(names := row["_SubBossInfoArray"].split("\n")) == len(set(names))
                    for row in data.rows if row["_SubBossInfoArray"]
                ))

    def test_msdata_only_examples_fill_only_confirmed_columns(self):
        data = build_mission_workbook_data(self.catalog, self.text_source, 1)
        rows = {row["_MissionId"]: row for row in data.rows}
        self.assertEqual(rows["MISSION_001040"]["_TitleMsg"], "Back to Camp")
        self.assertTrue(rows["MISSION_001040"]["_DetailMsg"])
        self.assertTrue(all(
            value == "" for column, value in rows["MISSION_400063"].items()
            if column != "_MissionId"
        ))

    def test_workbook_merges_mission_fields_and_centers_them(self):
        data = build_mission_workbook_data(self.catalog, self.text_source, 1)
        BASE_DIR.joinpath(".agents").mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=BASE_DIR / ".agents") as temp_dir:
            path = write_mission_workbook(Path(temp_dir) / "MissionData.xlsx", data)
            workbook = load_workbook(path)
            try:
                sheet = workbook["MissionData"]
                self.assertEqual([cell.value for cell in sheet[1]], list(MISSION_COLUMNS))
                self.assertEqual(sheet.max_row, len(data.rows) + 1)
                first, last = next(group for group in data.groups if group[1] > group[0])
                for column in ("_MissionId", "_DetailMsg", "_SubBossInfoArray"):
                    letter = get_column_letter(MISSION_COLUMNS.index(column) + 1)
                    self.assertIn(f"{letter}{first}:{letter}{last}", {
                        str(merged) for merged in sheet.merged_cells.ranges
                    })
                    self.assertEqual(sheet[f"{letter}{first}"].alignment.vertical, "center")
                target_letter = get_column_letter(MISSION_COLUMNS.index("MonsterName") + 1)
                self.assertNotIn(f"{target_letter}{first}:{target_letter}{last}", {
                    str(merged) for merged in sheet.merged_cells.ranges
                })
                condition_letter = get_column_letter(MISSION_COLUMNS.index("_TargetType") + 1)
                self.assertNotIn(f"{condition_letter}{first}:{condition_letter}{last}", {
                    str(merged) for merged in sheet.merged_cells.ranges
                })
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
