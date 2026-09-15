from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
import json
from pathlib import Path
import re
import unicodedata
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.hyperlink import Hyperlink

from config import (
    ENEMY_ACTION_WORKBOOK,
    SUPPORT_FILES,
    ZH_HANS_LANGUAGE_ID,
)
from src.data.text_db import TextDB, TextSource
from src.data.user3 import load_user3_table
from src.utils.log import file_size, info


ENEMY_ROOT_RELATIVE = Path("STM/GameDesign/Enemy")
SHELL_CREATOR_GLOB = "Em????/??/Shell/*ShellCreatorInfo*.user.3.json"
SHELL_CREATOR_ROOT_TYPE = "ace.user_data.ShellCreatorInfoData"
SHELL_CREATOR_ROW_TYPE = "ace.user_data.ShellCreatorInfoData.ShellCreatorInfo"
INDEX_SHEET_NAME = "怪物索引"
INDEX_HEADERS = ("怪物EM编号", "怪物名称")
ACTION_HEADERS = ("name", "comment")
INVALID_SHEET_CHARS_RE = re.compile(r"[\[\]:*?/\\]")
ENEMY_ID_RE = re.compile(r"^(EM(\d{4})_(\d{2}))_(\d+)$", re.IGNORECASE)
EM_DIRECTORY_RE = re.compile(r"^Em(\d{4})$", re.IGNORECASE)
VARIANT_DIRECTORY_RE = re.compile(r"^(\d{2})$")

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10, color="1F1F1F")
LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")
HEADER_BORDER = Border(
    right=Side(style="thin", color="FFFFFF"),
    bottom=Side(style="thin", color="FFFFFF"),
)


@dataclass(frozen=True, slots=True)
class EnemyActionSheet:
    enemy_id: str
    enemy_name: str
    sheet_name: str
    rows: tuple[tuple[str, str], ...]
    source_path: Path
    raw_rows: int
    duplicate_rows: int
    blank_rows: int


@dataclass(frozen=True, slots=True)
class EnemyActionCatalog:
    sheets: tuple[EnemyActionSheet, ...]

    @property
    def raw_rows(self) -> int:
        return sum(sheet.raw_rows for sheet in self.sheets)

    @property
    def displayed_rows(self) -> int:
        return sum(len(sheet.rows) for sheet in self.sheets)

    @property
    def duplicate_rows(self) -> int:
        return sum(sheet.duplicate_rows for sheet in self.sheets)

    @property
    def blank_rows(self) -> int:
        return sum(sheet.blank_rows for sheet in self.sheets)


def export_enemy_action_workbook(
    output_dir: Path,
    natives_dir: Path,
    text_source: TextSource,
) -> Path:
    info("    Loading enemy ShellCreatorInfo action names")
    catalog = load_enemy_action_catalog(
        natives_dir,
        text_source.build(ZH_HANS_LANGUAGE_ID),
    )
    path = write_enemy_action_workbook(
        output_dir / ENEMY_ACTION_WORKBOOK,
        catalog,
    )
    info(
        f"    Saved workbook: {path} ({file_size(path)}, "
        f"{len(catalog.sheets)} enemy sheet(s), "
        f"{catalog.displayed_rows} displayed row(s), "
        f"{catalog.duplicate_rows} duplicate row(s) removed, "
        f"{catalog.blank_rows} blank row(s) skipped)"
    )
    return path


def load_enemy_action_catalog(
    natives_dir: Path,
    text_db: TextDB,
) -> EnemyActionCatalog:
    natives_dir = Path(natives_dir)
    enemy_root = natives_dir / ENEMY_ROOT_RELATIVE
    sources = list(enemy_root.glob(SHELL_CREATOR_GLOB))
    if not sources:
        raise FileNotFoundError(f"No enemy ShellCreatorInfo files found under: {enemy_root}")

    candidates_by_prefix = _enemy_candidates(natives_dir, text_db)
    source_by_prefix: dict[str, Path] = {}
    sheets = []

    for source in sorted(sources, key=lambda path: _source_key(path, enemy_root)):
        prefix = _source_prefix(source, enemy_root)
        previous_source = source_by_prefix.get(prefix)
        if previous_source is not None:
            raise ValueError(
                f"Multiple ShellCreatorInfo files map to {prefix}: "
                f"{previous_source} and {source}"
            )
        source_by_prefix[prefix] = source

        candidates = candidates_by_prefix.get(prefix, [])
        if len(candidates) != 1:
            candidate_ids = ", ".join(enemy_id for enemy_id, _ in candidates) or "none"
            raise ValueError(
                f"ShellCreatorInfo path {source} must map to exactly one full enemy ID "
                f"for {prefix}; candidates: {candidate_ids}"
            )

        enemy_id, enemy_name = candidates[0]
        enemy_name = _required_text(enemy_name, "EnemyName", source)
        sheet_name = _enemy_sheet_name(enemy_name, enemy_id)
        records = _load_shell_creator_records(source)
        rows, duplicate_rows, blank_rows = _action_rows(records, source)
        sheets.append(
            EnemyActionSheet(
                enemy_id=enemy_id,
                enemy_name=enemy_name,
                sheet_name=sheet_name,
                rows=tuple(rows),
                source_path=source,
                raw_rows=len(records),
                duplicate_rows=duplicate_rows,
                blank_rows=blank_rows,
            )
        )

    sheets.sort(key=lambda sheet: _enemy_id_key(sheet.enemy_id))
    sheet_names = [INDEX_SHEET_NAME, *(sheet.sheet_name for sheet in sheets)]
    if len(sheet_names) != len(set(sheet_names)):
        raise ValueError("Enemy action workbook contains duplicate sheet names")
    return EnemyActionCatalog(tuple(sheets))


def write_enemy_action_workbook(
    path: Path,
    catalog: EnemyActionCatalog,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    index_sheet = workbook.active
    index_sheet.title = INDEX_SHEET_NAME
    index_sheet.append(INDEX_HEADERS)

    for sheet_data in catalog.sheets:
        sheet = workbook.create_sheet(sheet_data.sheet_name)
        sheet.append(ACTION_HEADERS)
        for name, comment in sheet_data.rows:
            sheet.append((_literal_text(name), _literal_text(comment)))

        index_sheet.append((sheet_data.enemy_id, sheet_data.enemy_name))
        name_cell = index_sheet.cell(index_sheet.max_row, 2)
        destination = sheet_data.sheet_name.replace("'", "''")
        name_cell.hyperlink = Hyperlink(
            ref=name_cell.coordinate,
            location=f"'{destination}'!A1",
            display=sheet_data.enemy_name,
        )
        name_cell.font = LINK_FONT

    _style_index_sheet(index_sheet)
    for sheet in workbook.worksheets[1:]:
        _style_action_sheet(sheet)

    workbook.active = 0
    workbook.save(path)
    return path


def _enemy_candidates(
    natives_dir: Path,
    text_db: TextDB,
) -> dict[str, list[tuple[str, str]]]:
    enemy_path = Path(natives_dir) / SUPPORT_FILES["enemy"]
    rows = load_user3_table(enemy_path, text_db)
    candidates: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for row in rows:
        enemy_id = row.get("enemyId")
        if not isinstance(enemy_id, str):
            continue
        match = ENEMY_ID_RE.fullmatch(enemy_id)
        if not match:
            continue
        candidates[match.group(1).upper()].append(
            (enemy_id.upper(), row.get("EnemyName"))
        )
    return candidates


def _load_shell_creator_records(path: Path) -> list[dict[str, Any]]:
    with Path(path).open("r", encoding="utf-8") as file:
        data = json.load(file)

    if not isinstance(data, list) or len(data) != 1 or not isinstance(data[0], dict):
        raise ValueError(f"Invalid ShellCreatorInfo root in {path}")
    payload = data[0].get(SHELL_CREATOR_ROOT_TYPE)
    if not isinstance(payload, dict):
        raise ValueError(f"Missing {SHELL_CREATOR_ROOT_TYPE} in {path}")
    records = payload.get("_ShellCreatorInfos")
    if not isinstance(records, list):
        raise ValueError(f"Missing _ShellCreatorInfos array in {path}")

    result = []
    for index, wrapper in enumerate(records):
        if not isinstance(wrapper, dict):
            raise ValueError(f"Invalid ShellCreatorInfo row {index} in {path}")
        record = wrapper.get(SHELL_CREATOR_ROW_TYPE)
        if not isinstance(record, dict):
            raise ValueError(
                f"Missing {SHELL_CREATOR_ROW_TYPE} at row {index} in {path}"
            )
        if "_Name" not in record or "_Comment" not in record:
            raise ValueError(f"Missing _Name or _Comment at row {index} in {path}")
        result.append(record)
    return result


def _action_rows(
    records: list[dict[str, Any]],
    source: Path,
) -> tuple[list[tuple[str, str]], int, int]:
    rows = []
    seen = set()
    duplicate_rows = 0
    blank_rows = 0

    for index, record in enumerate(records):
        name = _optional_text(record.get("_Name"), "_Name", source, index)
        comment = _optional_text(record.get("_Comment"), "_Comment", source, index)
        if not name and not comment:
            blank_rows += 1
            continue
        key = (name, comment)
        if key in seen:
            duplicate_rows += 1
            continue
        seen.add(key)
        rows.append(key)
    return rows, duplicate_rows, blank_rows


def _source_prefix(path: Path, enemy_root: Path) -> str:
    relative = Path(path).relative_to(enemy_root)
    if len(relative.parts) != 4 or relative.parts[2].lower() != "shell":
        raise ValueError(f"Unexpected enemy ShellCreatorInfo path: {path}")
    em_match = EM_DIRECTORY_RE.fullmatch(relative.parts[0])
    variant_match = VARIANT_DIRECTORY_RE.fullmatch(relative.parts[1])
    if not em_match or not variant_match:
        raise ValueError(f"Unexpected enemy ShellCreatorInfo path: {path}")
    return f"EM{em_match.group(1)}_{variant_match.group(1)}".upper()


def _source_key(path: Path, enemy_root: Path) -> tuple[int, int, str]:
    prefix = _source_prefix(path, enemy_root)
    match = re.fullmatch(r"EM(\d{4})_(\d{2})", prefix)
    if not match:
        raise ValueError(f"Invalid enemy prefix: {prefix}")
    return int(match.group(1)), int(match.group(2)), path.name.lower()


def _enemy_id_key(enemy_id: str) -> tuple[int, int, int]:
    match = ENEMY_ID_RE.fullmatch(enemy_id)
    if not match:
        raise ValueError(f"Invalid full enemy ID: {enemy_id}")
    return int(match.group(2)), int(match.group(3)), int(match.group(4))


def _enemy_sheet_name(enemy_name: str, enemy_id: str) -> str:
    name = f"{enemy_name}（{enemy_id}）"
    if len(name) > 31:
        raise ValueError(f"Enemy sheet name exceeds Excel's 31-character limit: {name}")
    if INVALID_SHEET_CHARS_RE.search(name):
        raise ValueError(f"Enemy sheet name contains an invalid Excel character: {name}")
    return name


def _required_text(value: Any, field: str, source: Path) -> str:
    if not isinstance(value, str):
        raise ValueError(f"{field} must be text for {source}")
    text = value.strip()
    if not text:
        raise ValueError(f"{field} is blank for {source}")
    return text


def _optional_text(
    value: Any,
    field: str,
    source: Path,
    index: int,
) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        raise ValueError(f"{field} must be text or null at row {index} in {source}")
    return value.strip()


def _literal_text(value: str) -> str:
    return "'" + value if value.startswith("=") else value


def _style_index_sheet(sheet) -> None:
    _style_sheet(sheet)
    sheet.sheet_properties.tabColor = "1F4E78"
    sheet.column_dimensions["A"].width = 18.0
    sheet.column_dimensions["B"].width = min(
        36.0,
        max(14.0, max((_text_width(cell.value) for cell in sheet["B"]), default=0.0) + 2),
    )


def _style_action_sheet(sheet) -> None:
    _style_sheet(sheet)
    sheet.column_dimensions["A"].width = _fitted_width(sheet, 1, 14.0, 52.0)
    sheet.column_dimensions["B"].width = _fitted_width(sheet, 2, 14.0, 72.0)
    for row in sheet.iter_rows(min_row=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)


def _style_sheet(sheet) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 22.0
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = HEADER_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.hyperlink is None:
                cell.font = BODY_FONT
            cell.alignment = Alignment(horizontal="left", vertical="center")


def _fitted_width(sheet, column: int, minimum: float, maximum: float) -> float:
    width = max(
        (_text_width(sheet.cell(row, column).value) for row in range(1, sheet.max_row + 1)),
        default=0.0,
    )
    return min(maximum, max(minimum, width + 2))


def _text_width(value: Any) -> float:
    if value is None:
        return 0.0
    width = 0.0
    for character in str(value):
        if unicodedata.east_asian_width(character) in {"W", "F"}:
            width += 2.0
        elif character.isascii() and character.isalnum():
            width += 1.2
        else:
            width += 1.0
    return width
