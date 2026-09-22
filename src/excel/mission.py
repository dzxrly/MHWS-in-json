"""Layout for the localized MissionData workbook."""

from math import ceil
from pathlib import Path
from unicodedata import east_asian_width

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.converters.mission import MissionWorkbookData
from src.excel.writer import write_workbook


HEADER_COLORS = {
    "mission": ("23466D", "345A81"),
    "target": ("355272", "496784"),
    "difficulty": ("1E6571", "337884"),
    "status": ("375C83", "4A7195"),
    "multiplayer": ("416A55", "557D67"),
    "size": ("806333", "97784A"),
}
BODY_COLORS = {
    "mission": ("FFFFFF", "EAF1F7"),
    "target": ("F7FAFD", "E8F0F8"),
    "difficulty": ("F0F8F8", "E5F1F2"),
    "status": ("F3F6FA", "E7EFF6"),
    "multiplayer": ("F1F8F3", "E6F1E9"),
    "size": ("FCF8EF", "F5EEDC"),
}
GROUP_BORDER = Side(style="thin", color="A8BACB")
SECTION_BORDER = Side(style="medium", color="8DA6B6")
WRAPPED_COLUMNS = {"_TitleMsg", "MonsterName", "_DetailMsg", "_SubBossInfoArray"}
BODY_ALIGNMENTS = {
    wrapped: Alignment(horizontal="left", vertical="center", wrap_text=wrapped)
    for wrapped in (False, True)
}
COLUMN_WIDTHS = {
    "_MissionId": 20,
    "_Version": 19,
    "_QuestType": 15,
    "_QuestAttribute": 19,
    "_QuestLv": 10,
    "_TitleMsg": 37,
    "_TargetType": 23,
    "_LegendaryID": 15,
    "MonsterName": 23,
    "_DetailMsg": 75,
    "_SubBossInfoArray": 28,
    "_DifficultyRankId": 37,
    "_MultiTableId": 37,
    "_RewardRank": 20,
    "_IsUseRandomSize": 18,
}


def write_mission_workbook(path: Path, data: MissionWorkbookData) -> Path:
    return write_workbook(
        path,
        {"MissionData": data.rows},
        80.0,
        lambda workbook: style_mission_workbook(workbook, data),
    )


def style_mission_workbook(workbook, data: MissionWorkbookData) -> None:
    sheet = workbook["MissionData"]
    sheet.insert_rows(1)
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 80
    sheet.freeze_panes = "I3"
    sheet.auto_filter.ref = None  # Excel cannot sort merged mission groups reliably.
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 42

    group_starts = set()
    group_colors = {}
    for index, header in enumerate(data.headers, start=1):
        if index == 1 or header.section != data.headers[index - 2].section or (
            header.bottom is not None and header.top != data.headers[index - 2].top
        ):
            group_starts.add(index)
        if header.bottom is not None and (header.section, header.top) not in group_colors:
            related = sum(
                1 for section, _ in group_colors if section == header.section
            )
            group_colors[(header.section, header.top)] = related % 2
        top_color, bottom_color = HEADER_COLORS[header.section]
        if header.bottom is not None and group_colors[(header.section, header.top)] % 2:
            top_color, bottom_color = bottom_color, top_color
        top = sheet.cell(1, index)
        bottom = sheet.cell(2, index)
        top.value = header.top
        bottom.value = header.bottom
        for cell, color in ((top, top_color), (bottom, bottom_color)):
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=SECTION_BORDER if index in group_starts else Side())
        letter = get_column_letter(index)
        width = COLUMN_WIDTHS.get(header.key, 18 if header.bottom else 14)
        sheet.column_dimensions[letter].width = width

    for index, header in enumerate(data.headers, start=1):
        if header.bottom is None:
            sheet.merge_cells(start_row=1, start_column=index, end_row=2, end_column=index)
        elif index in group_starts:
            end = index
            while end < len(data.headers) and (
                data.headers[end].section, data.headers[end].top
            ) == (header.section, header.top):
                end += 1
            if end > index:
                sheet.merge_cells(start_row=1, start_column=index, end_row=1, end_column=end)

    for group_index, (first_row, last_row) in enumerate(data.groups):
        row_count = last_row - first_row + 1
        first = data.rows[first_row - 3]
        needed_lines = max(
            _wrapped_line_count(first[column], COLUMN_WIDTHS[column] - 5)
            for column in ("_TitleMsg", "_DetailMsg", "_SubBossInfoArray")
        )
        height = min(120, max(24, ceil(needed_lines * 15 / row_count)))
        for row_index in range(first_row, last_row + 1):
            sheet.row_dimensions[row_index].height = height
            for column_index, header in enumerate(data.headers, start=1):
                column = header.key
                cell = sheet.cell(row_index, column_index)
                cell.fill = PatternFill(
                    "solid", fgColor=BODY_COLORS[header.section][group_index % 2]
                )
                cell.border = Border(
                    left=SECTION_BORDER if column_index in group_starts else Side(),
                    bottom=GROUP_BORDER if row_index == last_row else Side(),
                )
                cell.alignment = BODY_ALIGNMENTS[column in WRAPPED_COLUMNS]
        if row_count > 1:
            for column_index, header in enumerate(data.headers, start=1):
                if header.section == "mission":
                    sheet.merge_cells(
                        start_row=first_row,
                        start_column=column_index,
                        end_row=last_row,
                        end_column=column_index,
                    )


def _wrapped_line_count(value: str, width: int) -> int:
    return sum(
        max(1, ceil(sum(2 if east_asian_width(char) in "WF" else 1 for char in line) / width))
        for line in value.split("\n")
    )
