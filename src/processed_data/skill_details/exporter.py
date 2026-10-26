"""SkillDetails.xlsx: SkillCommonData and SkillData joined, with native effect links.

The links come from ``data/skill_links/`` (``index.json`` plus one shard per
skill category), written offline by ``python -m sdk.skill_effect_exporter
extract``. This builder never reads the EXE or il2cpp dump; rerun the
extractor after updating MHWS-in-json to a new game version.
"""

from __future__ import annotations

import json
import re
from copy import copy
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill

from config import (
    JAPANESE_LANGUAGE_ID,
    MAX_COLUMN_WIDTH,
    SUPPORT_FILES,
    ZH_HANS_LANGUAGE_ID,
)
from src.shared.excel.styles import StylePalette
from src.shared.excel.writer import write_workbook
from src.shared.source.repository import SourceRepository
from src.shared.text.catalog import TextSource

from .labels import DIMENSION_NAMES, ENUMS, WEAPON_CODES, effect_name, enum_label

OUTPUT_NAME = "SkillDetails.xlsx"
LINKS_DIR = Path(__file__).resolve().parent / "data" / "skill_links"
# Reviewed meaning of each shown effect, written by the SDK from its skill
# modules into the link JSON ("meanings": skill ID -> effect key (see
# effect_links) -> fields). "效果" replaces the generated name; the columns
# below are shown; "依据" (the code it was read from) stays in the JSON;
# "情形" names the situations a value depends on (see value_situations).
MEANING_COLUMNS = ("作用于", "计算", "条件", "单位")
MEANING_FIELDS = ("效果", *MEANING_COLUMNS, "依据", "情形")
SKILL_DATA = "STM/GameDesign/Common/Equip/SkillData.user.3.json"
SCHEMA_VERSION = 1
# Remarks at the end of in-game skill descriptions that the table leaves out.
DROPPED_NOTES = (
    "（骑乘鹭鹰龙时不会发动）",
    "（武器属性值的强化存在上限）",
    "（效果结束后，需要经过一定时间才能再次发动）",
    "（不会超过上限）",
    "（上升效果因武器种类而异）",
    "（部分攻击除外）",
    "（30分钟后达到增幅上限）",
    "（对甲虫种等部分怪物无效）",
    "（部分怪物除外）",
    "（满足对应条件后，效果可重复发动）",
)
HEADER_FILL = "305496"
# Rows alternate per skill; the value column uses a warmer tint of the same band.
BAND_FILLS = ("FFFFFF", "EAF1FB")
VALUE_FILLS = ("FFF4D6", "FBE8B5")
VALUE_COLUMN = "数值"
DIMENSION_EXCLUDED = {"INVALID", "MAX"}
OWN_SOURCES = {"PlayerSkillParam", "SkillData"}
OTHER_SOURCES = {"PlayerMealSkillParam"}
# A condition that another effect is running: a meal effect
# ("_MealEffect._IsEffectActive != 0", "_MealEffect...Timer > 0.0"), a hunting
# horn melody ("_Wp05MusicSkill._SkillTimer+0x90 > 0.0") or a weapon state flag
# tested bit by bit ("位 ..._HunterWeaponFlag..."; the bounds check before it,
# "... >= ...", is the flag not being set).
OTHER_ACTIVE = re.compile(
    r"_MealEffect\..*(!= 0|> 0)|^0 < .*_MealEffect\."
    r"|_Wp05MusicSkill\..* > 0\.0$|^位 .*_HunterWeaponFlag\."
)
# A random check against a percentage: "80 > (xorshift of the random state)".
RANDOM_CHECK = re.compile(r"^(\d+(?:\.\d+)?) > .*全局量 0x[0-9a-f]+ shl 15")
READING_NOTES = [
    "技能效果表每行是一个效果在某个技能等级的数值；“倍率”是乘算系数（1.05 即提高 5%），“加算”是直接加上的数值，“概率”是百分比",
    "“适用”为空表示对所有武器、弹种和动作都相同；列出时只对这些武器、弹种或动作生效",
    "同一等级在不同情形下（如连续命中次数、搬运时是否在爬墙）取不同数值时分行列出，情形写在“条件”列开头；只在其他效果（饭技、狩猎笛旋律、武器自身状态）生效时才出现的数值不列出；随机判定的效果写作发动概率（如 80%）",
    "仍写作“a / b”的数值，是代码中找不到区分条件的几种取值",
    "时间、距离等数值保留游戏内部单位，未换算",
    "没有把握翻译的动作、射击类型保留游戏内部名称（如 QUICK_SHOT、EX_FIELD），效果名称中的英文同理",
    "只列出从游戏参数读取到的数值；只有发动条件、没有数值的技能在“技能”表中标为“未提取到”",
    "“代码名称”是游戏中读取该数值的函数名，便于核对",
    "“作用于”“计算”“条件”“单位”来自数值含义表（SDK 中每个技能的模块 sdk/skill_effect_exporter/skills/<技能ID>.py），是对照代码中数值的用途写出的；标注“游戏说明”的内容取自技能说明，标注“推测”“未核实”的内容尚无代码证据，空白表示该效果还没有整理",
    "效果以“写入”开头的行，是技能发动等处理写入运行时字段的数值（如持续时间），名称保留代码中的字段名；同一字段按情形写入不同数据时，括号中注明数据字段",
    "技能自身的取值函数没有读到、只在命中处理等函数中读到的数据（如弹道强化 Lv3 的攻击力加算）也单独列出，按数据字段命名",
]


def load_links(directory: Path = LINKS_DIR) -> dict:
    """Reassemble the link document from index.json and its category shards."""
    directory = Path(directory)
    document = json.loads((directory / "index.json").read_text(encoding="utf-8"))
    if document.get("schemaVersion") != SCHEMA_VERSION:
        raise ValueError(f"技能链路索引版本不是 {SCHEMA_VERSION}：{directory}")
    listed = {entry["file"] for entry in document.pop("shards")}
    present = {path.name for path in directory.glob("*.json")} - {"index.json"}
    if listed != present:
        raise ValueError(
            f"技能链路分片与索引不一致：缺少 {sorted(listed - present)}，多出 {sorted(present - listed)}"
        )
    links = []
    for name in sorted(listed):
        shard = json.loads((directory / name).read_text(encoding="utf-8"))
        if shard.get("schemaVersion") != SCHEMA_VERSION:
            raise ValueError(f"技能链路分片的格式与索引不一致：{name}")
        links.extend(
            unpack_link(link, shard["strings"], shard["constants"])
            for link in shard["links"]
        )
    document["links"] = links
    return document


def unpack_link(link: dict, strings: list, constants: list) -> dict:
    """Replace string and constant table positions by their values."""

    def path(item):
        item = dict(item)
        for field in ("conditions", "compared", "sources", "opaque"):
            if field in item:
                item[field] = [strings[i] for i in item[field]]
        for field in ("expression", "detail"):
            if field in item:
                item[field] = strings[item[field]]
        if "constants" in item:
            item["constants"] = [constants[i] for i in item["constants"]]
        if "writes" in item:
            item["writes"] = [
                dict(
                    w,
                    target=strings[w["target"]],
                    **(
                        {"sources": [strings[i] for i in w["sources"]]}
                        if "sources" in w
                        else {}
                    ),
                    **(
                        {"constants": [constants[i] for i in w["constants"]]}
                        if "constants" in w
                        else {}
                    ),
                )
                for w in item["writes"]
            ]
        return item

    link = dict(link)
    for part in ("results", "usage"):
        if part in link:
            groups = [
                dict(group, paths=[path(p) for p in group["paths"]])
                for group in link[part]["groups"]
            ]
            link[part] = dict(link[part], groups=groups)
    return link


def cartesian_blocks(points: list[tuple]) -> list[list[set]]:
    """Split distinct points into blocks whose per-dimension value sets multiply out
    to exactly the block's points, so a row never implies combinations it lacks."""
    if len(points[0]) == 1:
        return [[{p[0] for p in points}]]
    lasts: dict[tuple, set] = {}
    for p in points:
        lasts.setdefault(p[:-1], set()).add(p[-1])
    by_last: dict[frozenset, list[tuple]] = {}
    for prefix, values in lasts.items():
        by_last.setdefault(frozenset(values), []).append(prefix)
    return [
        block + [set(values)]
        for values, prefixes in by_last.items()
        for block in cartesian_blocks(prefixes)
    ]


def sourced(path: dict) -> bool:
    """A finished path whose value was read from game data, not a code fallback."""
    return path["status"] == "ok" and "value" in path and bool(path.get("sources"))


def own_values(paths: list[dict], skills: set[str] | None = None) -> list[dict]:
    """The data-sourced paths that show the skill's own effect.

    A path reading only PlayerGlobalParam is the value without the skill; one
    that reads PlayerMealSkillParam, or runs with a meal effect active (the meal
    raises the skill level), stacks a meal effect on top. Both are dropped when
    a path reads the skill's own parameters alone. With skills given, a path
    reading another skill's SkillData row stacks that skill's effect (锁刃龙之饥饿
    reading 暗黑骑士之证's heal); like a path that needs another effect (meal,
    melody, weapon state) running, it is dropped even when nothing else is left."""
    paths = [p for p in paths if sourced(p)]

    def rows(path):
        return {
            m.group(1)
            for source in path["sources"]
            if (m := re.match(r"SkillData\[(\S+) ", source))
        }

    clean = [
        p
        for p in paths
        if not any(OTHER_ACTIVE.search(c) for c in p.get("conditions", []))
        and not (skills and rows(p) - skills)
    ]
    own = []
    for path in clean:
        roots = {
            re.split(r"[.\[]", source, maxsplit=1)[0] for source in path["sources"]
        }
        if roots & OWN_SOURCES and not roots & OTHER_SOURCES:
            own.append(path)
    return own or clean


def bare_method(function: dict) -> str:
    return re.sub(r"\d+$", "", function.get("method") or "")


def link_skills(link: dict) -> set[str]:
    """The skill a link is shown for and the skill read in the code (an open
    skill of a group): their SkillData rows are the link's own data."""
    return {link["skill"], link["via"]}


def source_keys(link: dict) -> set[str]:
    """The skill's own data items a link's shown values come from, without the
    level: "PlayerSkillParam._Field" or "SkillData[HunterSkill_X]._value[i]"."""
    keys = set()
    for group in link["results"]["groups"]:
        for path in own_values(group["paths"], link_skills(link)):
            for source in path["sources"]:
                slot = re.match(r"SkillData\[(\S+) [^\]]*\]\.(_value\[\d+\])", source)
                if slot:
                    keys.add(f"SkillData[{slot.group(1)}].{slot.group(2)}")
                elif re.split(r"[.\[]", source, maxsplit=1)[0] in OWN_SOURCES:
                    keys.add(".".join(re.split(r"[.\[]", source)[:2]))
    return keys


def effect_links(links: dict) -> dict[str, list[tuple[dict, str, str]]]:
    """Per skill, the links shown in the simplified sheet with their effect names
    and keys. The key names the entry in the meanings independently of the
    display name: the function ("getSkillAttackRate"), the function with the
    data items it is shown for ("getOnHitAttackPowerAdd：PlayerSkillParam._X"),
    or the function, the runtime field written and its data
    ("update→_Info._Timer：_X").

    Only links with data-sourced values count. Level getters are dropped when the
    skill has anything else, and the skill's own getters (getSkill*/isSkill*)
    win over callers that combine them with other state. A caller still shows
    when it reads data items of the skill that no shown getter reads (a bonus
    computed inside a hit handler); it is named after those items. Such an
    item counts only when no other skill's links use it: a function summing
    many skills evaluates the others' data in every scenario, so those items
    are not this skill's. Items of a link the skill's module declares (uses
    kind "module", with reviewed evidence) are the skill's even when the scan
    also ties the function to another skill."""
    functions = links["functions"]
    chosen: dict[str, list[dict]] = {}
    users: dict[str, set] = {}  # data item -> skills whose links use it
    declared: dict[str, set] = {}  # skill -> data items of its declared links
    for link in links["links"]:
        results = link.get("results")
        if results and any(
            own_values(group["paths"], link_skills(link)) for group in results["groups"]
        ):
            chosen.setdefault(link["skill"], []).append(link)
        module = bool(link["uses"]) and link["uses"][0]["kind"] == "module"
        for item in [link, *written_links(link)]:
            if item.get("results"):
                keys = source_keys(item)
                for key in keys:
                    users.setdefault(key, set()).add(link["skill"])
                if module:
                    declared.setdefault(link["skill"], set()).update(keys)

    def own_extra(keys, skill, covered):
        return sorted(
            k
            for k in keys - covered
            if users[k] == {skill} or k in declared.get(skill, ())
        )

    named: dict[str, list[tuple[dict, str, str]]] = {}
    shown_items: dict[str, set] = {}  # skill -> data items already shown
    for skill, own in chosen.items():
        method = {id(l): bare_method(functions[l["function"]]) for l in own}
        effects = [
            l for l in own if not re.search(r"(Lv|Level)$", method[id(l)])
        ] or own
        direct = [
            l
            for l in effects
            if method[id(l)].lower().startswith(("getskill", "isskill"))
        ]
        shown = direct or effects
        named[skill] = [
            (
                l,
                effect_name(functions[l["function"]]["method"]),
                bare_method(functions[l["function"]]),
            )
            for l in shown
        ]
        covered = set().union(*(source_keys(l) for l in shown))
        for link in effects:
            if any(link is l for l in shown):
                continue
            extra = own_extra(source_keys(link), skill, covered)
            if extra:
                key = f"{bare_method(functions[link['function']])}：{'、'.join(extra)}"
                named[skill].append((link, item_effect_name(extra), key))
                covered.update(extra)
        shown_items[skill] = covered
    # Values a handler writes into runtime state (a duration stored when the
    # skill activates), from data items nothing above shows.
    writes: dict[tuple[str, str], list[tuple[dict, list[str]]]] = {}
    for link in links["links"]:
        for written in written_links(link):
            covered = shown_items.setdefault(link["skill"], set())
            extra = own_extra(source_keys(written), link["skill"], covered)
            if extra:
                # Named by the runtime field written, kept as in the code.
                target = ".".join(written["written"].split(".")[-2:])
                writes.setdefault((link["skill"], target), []).append((written, extra))
                covered.update(extra)
    for (skill, target), entries in writes.items():
        for written, extra in entries:
            # Different data written to one field (a timer set from one of two
            # parameters by situation): each row names its data.
            fields = "、".join(key.split(".", 1)[1] for key in extra)
            name = (
                f"写入 {target}" if len(entries) == 1 else f"写入 {target}（{fields}）"
            )
            key = f"{bare_method(functions[written['function']])}→{target}：{fields}"
            named.setdefault(skill, []).append((written, name, key))
    return named


def written_links(link: dict) -> list[dict]:
    """One pseudo-link per runtime field the link's paths write a data-sourced
    number to; its paths return that number under the same conditions. Only
    fields of a known declared type count: without it the written bits cannot
    be read as a number (a float stored to an unnamed offset)."""
    targets: dict[str, list[dict]] = {}
    for group in (link.get("results") or {}).get("groups", []):
        per_target: dict[str, list[dict]] = {}
        for path in group["paths"]:
            if path["status"] != "ok":
                continue
            for write in path.get("writes", []):
                value = write["value"]
                if (
                    isinstance(value, str)
                    or not write.get("sources")
                    or not write.get("type")
                    # a flag set from the skill level, not a parameter value
                    or all(s.endswith("._SkillLv") for s in write["sources"])
                ):
                    continue
                per_target.setdefault(write["target"], []).append(
                    dict(
                        status="ok",
                        value=value,
                        sources=write["sources"],
                        conditions=path.get("conditions", []),
                    )
                )
        for target, paths in per_target.items():
            targets.setdefault(target, []).append(
                dict(group, paths=paths, written=target)
            )
    return [
        dict(link, results=dict(link["results"], groups=groups), written=target)
        for target, groups in targets.items()
    ]


def item_effect_name(keys: list[str]) -> str:
    """Effect name of data items: a parameter field is named like a getter of
    the same words (per-weapon digits dropped), a SkillData slot by its index."""
    names = []
    for key in keys:
        root, item = key.split(".", 1)
        if root.startswith("SkillData"):
            slot = re.search(r"\d+", item).group()
            names.append(f"数值槽 {slot}")
        else:
            names.append(effect_name("getSkill" + re.sub(r"Wp\d\d", "", item)))
    return "、".join(dict.fromkeys(names))


def effect_points(
    link: dict, situations: dict[str, str] | None = None
) -> tuple[list[str], dict, dict]:
    """(dimensions, domains, {(point, situation): {level: values}}) of one link.

    A point holds one label per dimension (None when it does not apply). Boolean
    arguments are not dimensions because the dump does not name them; the weapon
    a data path belongs to (``_Wp11Data`` and so on) stands in for them when the
    function has no weapon dimension of its own. Where one level and point has
    several values, value_situations tells them apart; situations maps code
    texts to readable situation names (the meaning's "情形")."""
    results = link["results"]
    keys, booleans, domains = {}, [], {}
    for key, labels in results["argDomains"].items():
        raw = key.split(":", 1)[-1]
        if raw == "Boolean":
            booleans.append(key)
            continue
        name = DIMENSION_NAMES.get(raw, raw)
        keys[key] = name
        domains.setdefault(name, set()).update(set(labels) - DIMENSION_EXCLUDED)
    dims = list(dict.fromkeys(keys.values()))
    inferred = "武器" not in dims
    if inferred:
        dims.append("武器")
    entries: dict[tuple, list[tuple]] = {}  # (point, level) -> [(booleans, path)]
    skills = {link["skill"], link["via"]}
    for group in results["groups"]:
        paths = own_values(group["paths"], skills)
        if not paths:
            continue
        weapons = {
            WEAPON_CODES[code]
            for p in paths
            for source in p["sources"]
            for code in re.findall(r"Wp\d\d", source)
            if code in WEAPON_CODES
        }
        for case in group["cases"]:
            labels = dict.fromkeys(dims)
            for key, name in keys.items():
                labels[name] = labels[name] or case["args"].get(key)
            if inferred and len(weapons) == 1:
                labels["武器"] = next(iter(weapons))
            flags = tuple((key, case["args"].get(key)) for key in booleans)
            entries.setdefault((tuple(labels.values()), case["level"]), []).extend(
                (flags, p) for p in paths
            )
    points: dict[tuple, dict[int, set]] = {}
    for (point, level), items in entries.items():
        for situation, values in value_situations(items, situations or {}).items():
            levels = points.setdefault((point, situation), {})
            levels.setdefault(level, set()).update(values)
    return dims, domains, points


def value_situations(items: list[tuple], names: dict[str, str]) -> dict:
    """{situation or None: values} of the paths of one level and point.

    A random check ("80 > random") becomes its chance, "80%". Otherwise values
    differing by an unnamed boolean argument are split by it; then each value
    is named by the data field only its paths read (the weapon's table or a
    special action's table), else by the conditions only its paths share (a
    hit count). A value nothing tells apart stays unnamed (shown "a / b")."""
    paths = [p for _, p in items]
    values = {p["value"] for p in paths}
    if values <= {True, False}:
        chances = {
            m.group(1)
            for p in paths
            for c in p.get("conditions", [])
            if (m := RANDOM_CHECK.match(c))
        }
        if len(chances) == 1:
            return {None: {f"{float(chances.pop()):g}%"}}
    if len(values) == 1:
        return {None: values}
    by_flags: dict[tuple, list[dict]] = {}
    for flags, path in items:
        by_flags.setdefault(flags, []).append(path)
    if len(by_flags) == 1 or not all(
        len({p["value"] for p in ps}) == 1 for ps in by_flags.values()
    ):
        by_flags = {(): paths}
    out: dict = {}
    for flags, ps in by_flags.items():
        prefix = "、".join(
            names.get(
                f"参数{key.split(':')[0]}={value}",
                f"参数{key.split(':')[0]} 为 {value}",
            )
            for key, value in flags
        )
        for label, vals in condition_situations(ps, names).items():
            name = "；".join(x for x in (prefix, label) if x) or None
            out.setdefault(name, set()).update(vals)
    return out


def condition_situations(paths: list[dict], names: dict[str, str]) -> dict:
    by_value: dict = {}
    for path in paths:
        by_value.setdefault(path["value"], []).append(path)
    if len(by_value) == 1:
        return {None: set(by_value)}

    def common(ps, field):
        return set.intersection(*(set(field(p)) for p in ps))

    fields = {v: common(ps, source_fields) for v, ps in by_value.items()}
    conditions = {
        v: common(ps, lambda p: p.get("conditions", [])) for v, ps in by_value.items()
    }
    shared_fields = set.intersection(*fields.values())
    shared_conditions = set.intersection(*conditions.values())
    out: dict = {}
    for value in by_value:
        own_fields = sorted(fields[value] - shared_fields)
        own_conditions = sorted(conditions[value] - shared_conditions)
        parts = [
            next((t for k, t in names.items() if k in f), f"使用数据 {f}")
            for f in own_fields
        ]
        for c in own_conditions:
            # A name "" leaves a condition out (the "!= 2" of "== 1"); next to
            # a data field only named conditions are added (the hit count).
            name = next((t for k, t in names.items() if k in c), None)
            if name is None and not own_fields:
                name = readable(c)
            if name and name not in parts:
                parts.append(name)
        label = "；".join(parts)
        out.setdefault(label or None, set()).add(value)
    return out


def source_fields(path: dict) -> list[str]:
    """Data fields of a path's sources without indices: "_ContinuousAttackWp00Data",
    "SkillData 槽0"."""
    fields = []
    for source in path["sources"]:
        slot = re.match(r"SkillData\[[^\]]*\]\._value\[(\d+)\]", source)
        if slot:
            fields.append(f"SkillData 槽{slot.group(1)}")
        else:
            fields.append(re.split(r"[.\[]", source)[1])
    return fields


def readable(condition: str) -> str:
    condition = condition.replace("cHunterSkill._HunterSkillParamInfo.", "")
    condition = re.sub(r"<(\w+)>k__BackingField", r"\1", condition)
    condition = re.sub(r"(\w)\d{4,}\(\)", r"\1()", condition)
    return condition if len(condition) <= 80 else condition[:80] + "…"


def applies_texts(dims: list[str], domains: dict, block: list[set]) -> list[str]:
    """Texts such as "【重弩、轻弩】散弹": weapons in brackets, then ammo kinds or
    actions. Actions named after one weapon (``WP03_MIKIRI``) go under that
    weapon; a block mixing them with general actions becomes one text per
    weapon so that every text has at most one bracket."""
    weapons, parts, by_weapon = "", [], {}
    for name, labels in zip(dims, block):
        labels = labels - {None} - DIMENSION_EXCLUDED
        domain = domains.get(name, set())
        if not labels or (len(domain) > 1 and labels >= domain):
            continue
        order = list(ENUMS.get(name, {}))
        labels = sorted(
            labels, key=lambda v: (order.index(v) if v in order else len(order), v)
        )
        if name == "武器":
            weapons = "、".join(enum_label(name, label) for label in labels)
            continue
        plain = []
        for label in labels:
            prefix = re.match(r"WP(\d\d)_(.+)", label)
            weapon = prefix and WEAPON_CODES.get(f"Wp{prefix.group(1)}")
            if weapon:
                text = ENUMS.get(name, {}).get(label, prefix.group(2))
                by_weapon.setdefault(weapon, []).append(text)
            else:
                plain.append(enum_label(name, label))
        if plain:
            parts.append("、".join(plain))
    texts = []
    if parts or not by_weapon:
        texts.append((f"【{weapons}】" if weapons else "") + "；".join(parts))
    texts += [
        f"【{enum_label('武器', weapon)}】" + "、".join(actions)
        for weapon, actions in by_weapon.items()
    ]
    return texts


def description(text: str | None, dropped: tuple[str, ...] = ()) -> str:
    """In-game text without its hard line breaks and the given remarks."""
    text = re.sub(r"\r?\n", "", text or "")
    for note in dropped:
        text = text.replace(note, "")
    return text


def level_text(level: int) -> str:
    return f"Lv.{level}"


def cell_value(values: set):
    # Single-precision values: 6 significant digits drop the float32 tail.
    values = {float(f"{v:.6g}") if isinstance(v, float) else v for v in values}
    if len(values) == 1:
        (value,) = values
        return ("是" if value else "否") if isinstance(value, bool) else value
    return " / ".join(
        ("是" if v else "否") if isinstance(v, bool) else f"{v:g}"
        for v in sorted(values)
    )


def build_sheets(
    repository: SourceRepository,
    text_source: TextSource,
    links: dict,
) -> dict[str, list[dict]]:
    """技能效果: one row per skill level and effect; 技能: SkillCommonData x SkillData."""
    meanings = links["meanings"]
    zh = text_source.build(ZH_HANS_LANGUAGE_ID)
    ja = text_source.build(JAPANESE_LANGUAGE_ID)
    common_rows = [
        row
        for row in repository.table(SUPPORT_FILES["skill_common"])
        if row["skillId"] != "NONE"
    ]
    names = {
        row["skillId"]: (
            ja.get(row["skillName"]) or "",
            zh.get(row["skillName"]) or "",
            description(zh.get(row["skillExplain"]), DROPPED_NOTES),
        )
        for row in common_rows
    }
    levels_of = {}
    for row in repository.table(SKILL_DATA):
        if row["skillId"] != "NONE":
            levels_of.setdefault(row["skillId"], []).append(row)
    level_explain = {
        (skill_id, row["SkillLv"]): description(zh.get(row["skillExplain"]))
        for skill_id, rows in levels_of.items()
        for row in rows
    }
    skills = links["skills"]
    functions = links["functions"]
    chosen = effect_links(links)
    effects = []
    order = [row["skillId"] for row in common_rows]
    order += sorted(set(chosen) - set(order))
    for skill_id in order:
        ja_name, zh_name, explain = names.get(skill_id, ("", "", ""))
        rows = []
        for link, effect, key in chosen.get(skill_id, []):
            function = functions[link["function"]]
            meaning = meanings.get(skill_id, {}).get(key, {})
            dims, domains, points = effect_points(link, meaning.get("情形"))
            # Per level and situation, points with the same values share one row.
            by_level: dict[tuple, list[tuple]] = {}
            for (point, situation), levels in points.items():
                for level, values in levels.items():
                    key = (level, frozenset(values), situation)
                    by_level.setdefault(key, []).append(point)
            for (level, values, situation), members in by_level.items():
                condition = "；".join(
                    x for x in (situation, meaning.get("条件", "")) if x
                )
                for block in cartesian_blocks(members):
                    for applies in applies_texts(dims, domains, block):
                        rows.append(
                            {
                                "技能ID": skill_id,
                                "日文名称": ja_name,
                                "简体中文名称": zh_name,
                                "技能说明": explain,
                                "技能等级": level,
                                "数值": cell_value(set(values)),
                                "效果": meaning.get("效果") or effect,
                                "适用": applies,
                                "条件": condition,
                                "等级说明": level_explain.get((skill_id, level), ""),
                                **{
                                    c: meaning.get(c, "")
                                    for c in MEANING_COLUMNS
                                    if c != "条件"
                                },
                                "备注": "；".join(
                                    note
                                    for note, shown in (
                                        (
                                            "数值随情形不同，未找到区分条件",
                                            len(values) > 1,
                                        ),
                                        ("技能发动等处理写入的数值", "written" in link),
                                    )
                                    if shown
                                ),
                                "代码名称": bare_method(function),
                            }
                        )
        rows.sort(key=lambda r: (r["技能等级"], r["效果"], r["适用"], r["条件"]))
        for row in rows:
            row["技能等级"] = level_text(row["技能等级"])
        effects += rows

    counted = {row["技能ID"] for row in effects}
    notes = links["notes"]["skills"]
    overview = []
    for row in common_rows:
        skill_id = row["skillId"]
        ja_name, zh_name, explain = names[skill_id]
        info = skills.get(skill_id, {})
        for data in sorted(levels_of.get(skill_id, []), key=lambda r: r["SkillLv"]) or [
            None
        ]:
            overview.append(
                {
                    "技能ID": skill_id,
                    "类别": row["skillCategory"],
                    "日文名称": ja_name,
                    "简体中文名称": zh_name,
                    "技能说明": explain,
                    "技能等级": level_text(data["SkillLv"]) if data else "",
                    "等级说明": (
                        description(zh.get(data["skillExplain"])) if data else ""
                    ),
                    "打开技能": "、".join(
                        names.get(s, ("", s, ""))[1] or s
                        for s in info.get("openSkills", [])
                    ),
                    "效果数值": "见技能效果表" if skill_id in counted else "未提取到",
                    "说明": "\n".join(notes.get(skill_id, [])),
                }
            )

    remarks = [{"项目": "读法", "内容": text} for text in READING_NOTES]
    remarks += [
        {"项目": "通用说明", "内容": text} for text in links["notes"]["general"]
    ]
    return {"技能效果": effects, "技能": overview, "说明": remarks}


def export_skill_details(
    output_dir: Path, repository: SourceRepository, text_source: TextSource
) -> Path:
    links = load_links()
    sheets = build_sheets(repository, text_source, links)
    return write_workbook(
        Path(output_dir) / OUTPUT_NAME, sheets, MAX_COLUMN_WIDTH, format_workbook
    )


def format_workbook(workbook) -> None:
    """Header colours, one background band per skill, a tinted value column, frozen
    names and filters."""
    palette = StylePalette(workbook, "skill_details")
    align = Alignment(horizontal="left", vertical="center")
    bold = copy(palette.default_font)
    bold.b = True
    header_font = copy(bold)
    header_font.color = "FFFFFF"
    header = palette.get(
        "header", font=header_font, fill=solid(HEADER_FILL), alignment=align
    )
    for sheet in workbook.worksheets:
        titles = [cell.value for cell in sheet[1]]
        for cell in sheet[1]:
            cell.style = header
        sheet.auto_filter.ref = sheet.dimensions
        if "简体中文名称" in titles:
            sheet.freeze_panes = sheet.cell(2, titles.index("简体中文名称") + 2)
        else:
            sheet.freeze_panes = "A2"
        key = titles.index("技能ID") if "技能ID" in titles else None
        band, last = 0, None
        for row in sheet.iter_rows(min_row=2):
            if key is not None and row[key].value != last:
                band = 1 - band if last is not None else 0
                last = row[key].value
            for cell in row:
                value = titles[cell.column - 1] == VALUE_COLUMN
                fill = (VALUE_FILLS if value else BAND_FILLS)[band]
                cell.style = palette.get(
                    (fill, value),
                    font=bold if value else palette.default_font,
                    fill=solid(fill),
                    alignment=align,
                )


def solid(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def usage_evidence(links: dict, link: dict, key: str, limit: int = 10) -> list[str]:
    """Usage-tracing lines about one shown effect, for writing its meaning.

    The effect's own reads are named in other links' usage as 取值〈getter〉
    (a call of the getter) or 技能数值〈skill 槽k〉 (its SkillData slot); a
    caller or write row is also traced itself. Only when none of these
    appear, lines naming the skill's level (技能等级〈skill〉) are used, marked
    as such: the getter may be inlined where the level is read."""
    functions = links["functions"]
    method = re.split("[：→]", key)[0]
    labels = {f"取值〈{method}〉"}
    for group in link["results"]["groups"]:
        for path in own_values(group["paths"], link_skills(link)):
            for source in path["sources"]:
                slot = re.match(r"SkillData\[(\S+) Lv\d+\]\._value\[(\d+)\]", source)
                if slot:
                    labels.add(f"技能数值〈{slot.group(1)} 槽{slot.group(2)}〉")
    level = f"技能等级〈{link['via']}〉"

    def lines(wanted, own_function_only):
        found = []
        for other in links["links"]:
            usage = other.get("usage")
            if not usage or not wanted & set(usage["labels"]):
                continue
            if own_function_only and other["function"] != link["function"]:
                continue
            function = functions[other["function"]]
            where = (
                f"{(function['type'] or '').rsplit('.', 1)[-1]}.{bare_method(function)}"
            )
            for group in usage["groups"]:
                for path in group["paths"]:
                    texts = []
                    result = path.get("expression", path.get("value"))
                    if any(label in str(result) for label in wanted):
                        texts.append(f"返回 {result}")
                    texts += [
                        f"写入 {w['target']} = {w['value']}"
                        for w in path.get("writes", [])
                        if any(label in str(w["value"]) for label in wanted)
                    ]
                    texts += [
                        f"条件 {c}"
                        for c in path.get("conditions", [])
                        if any(label in c for label in wanted)
                    ]
                    for text in texts:
                        text = f"{where}：{text}"
                        if text not in found:
                            found.append(text)
        return found

    found = lines(labels, False) or lines({level}, True)
    if not found:
        found = [f"（按技能等级）{t}" for t in lines({level}, False)]
    return [t if len(t) <= 300 else t[:300] + "…" for t in found[:limit]]


def meanings_report(links: dict) -> str:
    """Markdown list of every shown effect: key, values, meaning, evidence; then
    meanings whose effect is no longer shown."""
    meanings = links["meanings"]
    chosen = effect_links(links)
    out = ["# 数值含义表核对清单", ""]
    shown = set()
    for skill in sorted(chosen):
        out.append(f"## {skill}")
        for link, effect, key in chosen[skill]:
            shown.add((skill, key))
            values = sorted(
                {
                    f"{v:g}" if isinstance(v, float) else str(v)
                    for group in link["results"]["groups"]
                    for p in own_values(group["paths"], link_skills(link))
                    for v in [p["value"]]
                }
            )
            meaning = meanings.get(skill, {}).get(key)
            out.append(f"- `{key}` {effect}；数值 {'、'.join(values[:8])}")
            if meaning:
                out.append(
                    "  - 含义："
                    + "；".join(
                        f"{f} {meaning[f]}" for f in MEANING_FIELDS if meaning.get(f)
                    )
                )
            else:
                out.append("  - 含义：（未填写）")
            out += [f"  - {line}" for line in usage_evidence(links, link, key)]
        out.append("")
    stale = [
        f"- {skill} `{key}`"
        for skill, entries in meanings.items()
        for key in entries
        if (skill, key) not in shown
    ]
    if stale:
        out += ["## 已不再显示的条目", "", *stale, ""]
    return "\n".join(out) + "\n"
