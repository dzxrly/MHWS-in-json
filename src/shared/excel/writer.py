from collections.abc import Callable
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from src.shared.excel.style import apply_rare_style, style_workbook
from src.shared.excel.cells import safe_cell, sheet_name, text_width
from src.shared.rarity import Rarity
from src.shared.tables import columns as table_columns


WorkbookFormatter = Callable[[Workbook], None]


def write_workbook(
    path: Path,
    sheets: dict[str, list[dict]],
    max_width: float,
    formatter: WorkbookFormatter | None = None,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    rarity_cells = []
    first = True
    for title, rows in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet()
        first = False
        sheet.title = sheet_name(title)
        columns = table_columns(rows)
        sheet.append(columns)
        widths = [max(6.0, text_width(column)) for column in columns]
        for row_index, row in enumerate(rows, start=2):
            values = []
            for index, column in enumerate(columns):
                value = row.get(column)
                if isinstance(value, Rarity):
                    values.append(value.level)
                    rarity_cells.append((sheet, row_index, index + 1, value.index))
                else:
                    value = safe_cell(value)
                    values.append(value)
                widths[index] = max(widths[index], text_width(value))
            sheet.append(values)
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = min(width, max_width)
    if not formatter:
        style_workbook(workbook, max_width)
    for sheet, row, column, rarity in rarity_cells:
        apply_rare_style(sheet.cell(row, column), rarity)
    if formatter:
        formatter(workbook)
    workbook.save(path)
    return path
