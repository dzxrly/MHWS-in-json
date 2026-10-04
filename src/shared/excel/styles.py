"""Reusable Excel style palettes, table alignment and rarity presentation."""

from copy import copy
from openpyxl.styles import NamedStyle
from openpyxl.styles.fonts import DEFAULT_FONT
from openpyxl.styles import Alignment, Font, PatternFill


class StylePalette:
    def __init__(self, workbook, prefix: str):
        self.workbook = workbook
        self.prefix = prefix
        self.styles = {}
        self.default_font = copy(DEFAULT_FONT)

    def get(self, key, **components) -> str:
        if key not in self.styles:
            components.setdefault("font", self.default_font)
            style = NamedStyle(name=f"{self.prefix}_{len(self.styles)}", **components)
            self.workbook.add_named_style(style)
            # Assign by registered name: passing NamedStyle objects makes openpyxl
            # compare complete styles against its list for every cell assignment.
            self.styles[key] = style.name
        return self.styles[key]


RARE_COLORS = {
    0: "969696",
    1: "DEDEDE",
    2: "A4C43B",
    3: "47A33F",
    4: "5CAEBB",
    5: "575FD9",
    6: "9272E3",
    7: "C76D46",
    8: "B3436A",
    9: "2EC9E6",
    10: "F2C21D",
    11: "B4F5FF",
}


def style_workbook(workbook, max_width: float = 80.0) -> None:
    palette = StylePalette(workbook, "table")
    body = palette.get(
        "body", alignment=Alignment(horizontal="left", vertical="center")
    )
    wrapped = palette.get(
        "wrapped",
        alignment=Alignment(horizontal="left", vertical="center", wrap_text=True),
    )
    header = palette.get(
        "header",
        font=Font(bold=True),
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    for sheet in workbook.worksheets:
        wrapped_columns = {
            cell.column for cell in sheet[1] if cell.value in {"Explain", "RawExplain"}
        }
        for row in sheet.iter_rows():
            for cell in row:
                cell.style = (
                    header
                    if cell.row == 1
                    else wrapped if cell.column in wrapped_columns else body
                )


def apply_rare_style(cell, rare: int, bold: bool = False) -> None:
    color = _blend_with_white(RARE_COLORS.get(rare, "FFFFFF"), 0.5)
    cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    cell.font = Font(
        bold=bold,
        color="FFFFFF" if _brightness(color) < 64 else "000000",
    )


def _brightness(color: str) -> float:
    r, g, b = int(color[:2], 16), int(color[2:4], 16), int(color[4:6], 16)
    return (r * 299 + g * 587 + b * 114) / 1000


def _blend_with_white(color: str, opacity: float) -> str:
    rgb = [int(color[i : i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{int(c * opacity + 255 * (1 - opacity)):02X}" for c in rgb)
